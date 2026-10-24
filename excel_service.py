# app/services/import_service.py
"""
خدمة معالجة واستيراد ملفات Excel والبحث عن الصور والرفع إلى Salla
مشتركة بين لوحة التحكم ولوحة الإدارة.
"""

import asyncio
import io
import re
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
from fastapi import HTTPException

from config import BASE_DIR
import logging
def get_logger(n): return logging.getLogger(n)
# auth
from image_finder.downloader import download_image
from image_finder.finder import ProductImageFinder
from openserp_manager import acquire_for_search, release_after_search
# salla

logger = get_logger(__name__)

# تتبع التقدم اللحظي للوظائف في الذاكرة
_progress_state: Dict[str, Dict[str, Any]] = {}


def get_progress(job_id: str) -> Dict[str, Any]:
    """استرجاع حالة تقدم عملية البحث عن الصور."""
    return _progress_state.get(
        job_id,
        {
            "processed": 0,
            "total": 0,
            "succeeded": 0,
            "failed": 0,
            "current": "",
            "status": "pending",
        },
    )


# ═══════════════════════════════════════════════════════════
# تحليل وقراءة ملفات Excel و CSV
# ═══════════════════════════════════════════════════════════
def parse_excel_file(content: bytes, filename: str) -> Dict[str, Any]:
    fname = filename.lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls") or fname.endswith(".csv")):
        raise ValueError("الملف يجب أن يكون .xlsx أو .xls أو .csv")

    try:
        df = _read_excel_file(content, fname)
    except Exception as e:
        raise ValueError(f"فشل قراءة الملف: {type(e).__name__}: {str(e)}")

    if df.empty:
        raise ValueError("الملف فارغ")

    column_map = _detect_columns(df)
    if not column_map.get("name"):
        raise ValueError(
            "لم يتم العثور على عمود الاسم. الأعمدة المتاحة: "
            + ", ".join(map(str, df.columns.tolist()))
        )

    products = []
    for idx, row in df.iterrows():
        product = {
            "row": int(idx) + 2,
            "no": _get_cell(row, column_map.get("no")),
            "name": _get_cell(row, column_map.get("name")),
            "barcode": _get_cell(row, column_map.get("barcode")),
            "sku": _get_cell(row, column_map.get("sku")),
            "price": _get_cell(row, column_map.get("price")),
            "sale_price": _get_cell(row, column_map.get("sale_price")),
            "quantity": _get_cell(row, column_map.get("quantity")),
            "description": _get_cell(row, column_map.get("description")),
            "category": _get_cell(row, column_map.get("category")),
            "type": _get_cell(row, column_map.get("type")),
            "product_type": _get_cell(row, column_map.get("product_type")),
            "image_url": _get_cell(row, column_map.get("image_url")),
            "image_alt": _get_cell(row, column_map.get("image_alt")),
            "status": _get_cell(row, column_map.get("status")),
            "weight": _get_cell(row, column_map.get("weight")),
            "weight_unit": _get_cell(row, column_map.get("weight_unit")),
            "brand": _get_cell(row, column_map.get("brand")),
            "requires_shipping": _get_cell(row, column_map.get("requires_shipping")),
            "taxable": _get_cell(row, column_map.get("taxable")),
            "cost": _get_cell(row, column_map.get("cost")),
            "calories": _get_cell(row, column_map.get("calories")),
            "mpn": _get_cell(row, column_map.get("mpn")),
            "gtin": _get_cell(row, column_map.get("gtin")),
            "promo_title": _get_cell(row, column_map.get("promo_title")),
            "pin": _get_cell(row, column_map.get("pin")),
            "max_per_customer": _get_cell(row, column_map.get("max_per_customer")),
            "hide_qty": _get_cell(row, column_map.get("hide_qty")),
            "image_on_order": _get_cell(row, column_map.get("image_on_order")),
        }
        if not product["name"]:
            continue
        products.append(product)

    if not products:
        raise ValueError("لا توجد منتجات صالحة في الملف")

    return {
        "success": True,
        "total": len(products),
        "columns_detected": column_map,
        "products": products,
    }


# ═══════════════════════════════════════════════════════════
# البحث عن الصور في الخلفية
# ═══════════════════════════════════════════════════════════
async def search_images_for_store(
    products: List[Dict[str, Any]],
    use_ai: bool = True,
    job_id: Optional[str] = None,
    store_id: Optional[str] = None,
) -> Dict[str, Any]:
    job_id = job_id or str(uuid.uuid4())

    if not products:
        raise ValueError("لم يتم استلام منتجات للبحث")

    products_input = [
        {"barcode": str(p.get("barcode") or ""), "name": str(p.get("name") or "")}
        for p in products
    ]

    _progress_state[job_id] = {
        "processed": 0,
        "total": len(products_input),
        "succeeded": 0,
        "failed": 0,
        "current": "🔧 تجهيز محرك البحث...",
        "status": "preparing",
    }

    # تشغيل OpenSERP عند الحاجة
    acquire_result = await acquire_for_search()
    if not acquire_result.get("ready"):
        _progress_state[job_id]["status"] = "failed"
        raise RuntimeError(
            f"تعذر تجهيز OpenSERP: {acquire_result.get('error', 'غير معروف')}"
        )

    _progress_state[job_id]["status"] = "processing"
    _progress_state[job_id]["current"] = ""

    start = time.time()

    def process_one(idx: int):
        product = products_input[idx]
        finder = ProductImageFinder(use_ai=use_ai)

        _progress_state[job_id]["current"] = product["name"]

        try:
            result = finder.find(
                product["barcode"],
                product_name=product["name"],
                force_refresh=False,
            )
        except Exception as e:
            result = {"success": False, "error": str(e)}

        original = products[idx]
        state = _progress_state[job_id]
        state["processed"] += 1
        if result.get("success"):
            state["succeeded"] += 1
        else:
            state["failed"] += 1

        return {
            "row": original.get("row"),
            "name": product["name"],
            "barcode": product["barcode"],
            "sku": original.get("sku") or product["barcode"],
            "price": original.get("price"),
            "quantity": original.get("quantity"),
            "description": original.get("description"),
            "category": original.get("category", ""),
            "product_type": original.get("product_type", ""),
            "status": original.get("status", "متاح"),
            "search_success": result.get("success", False),
            "search_score": result.get("score"),
            "image_url": result.get("image_url"),
            "candidates": (result.get("candidates") or [])[:5],
            "top_title": result.get("top_title", ""),
            "error": result.get("error"),
            "ai_vision": result.get("ai_vision"),
            "ai_identity": result.get("ai_identity"),
            "smart_search_used": result.get("smart_search_used", False),
            "decision": result.get("decision"),
        }

    MAX_WORKERS = 5
    loop = asyncio.get_event_loop()

    try:
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
            futures = [
                loop.run_in_executor(pool, process_one, i)
                for i in range(len(products_input))
            ]
            results_combined = await asyncio.gather(*futures)

        elapsed = time.time() - start
        state = _progress_state[job_id]
        state["status"] = "completed"
        state["current"] = ""

        summary = {
            "total": len(results_combined),
            "success": state["succeeded"],
            "failed": state["failed"],
            "elapsed": round(elapsed, 1),
        }

        rate = summary["total"] / elapsed if elapsed > 0 else 0
        logger.info(
            f"⚡ معالجة {summary['total']} منتج في {summary['elapsed']}s "
            f"({rate:.1f} منتج/ثانية) | نجح: {summary['success']} | فشل: {summary['failed']}"
        )

    finally:
        try:
            await release_after_search()
        except Exception as e:
            logger.exception(f"فشل إيقاف OpenSERP: {e}")

    return {
        "success": True,
        "job_id": job_id,
        "use_ai": use_ai,
        "summary": summary,
        "results": results_combined,
    }


# ═══════════════════════════════════════════════════════════
# رفع المنتجات المحددة إلى متجر Salla
# ═══════════════════════════════════════════════════════════
async def upload_products_to_salla(store_id: str, items: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not store_id:
        raise ValueError("يجب تحديد store_id لرفع المنتجات إلى Salla")
    if not items:
        raise ValueError("لم يتم تحديد أي منتج للرفع")

    token = await get_valid_access_token(store_id)
    client = SallaClient(token)

    results = []

    for idx, item in enumerate(items):
        name = (item.get("name") or "").strip()
        if not name:
            results.append(
                {
                    "index": idx,
                    "name": name,
                    "success": False,
                    "error": "اسم المنتج فارغ",
                }
            )
            continue

        try:
            price_value = _safe_float(item.get("price"))
            quantity_value = _safe_int(item.get("quantity"))
            barcode = _clean_str(item.get("barcode"))
            sku = _clean_str(item.get("sku")) or barcode
            weight, weight_unit = _extract_weight_from_name(name)

            if not price_value or price_value <= 0:
                price_value = 1.0

            payload = {
                "name": name[:200],
                "price": float(price_value),
                "product_type": "product",
                "quantity": int(quantity_value) if quantity_value else 0,
                "status": "sale",
                "require_shipping": True,
                "with_tax": True,
                "weight": float(weight),
                "weight_unit": weight_unit,
            }
            if sku:
                payload["sku"] = sku[:100]
            if item.get("description"):
                payload["description"] = str(item["description"])[:1000]

            # فحص وجود المنتج بالـ SKU
            existing_id = None
            if sku:
                search_resp = await client.get(
                    "/products", {"search": sku, "per_page": 50}
                )
                if search_resp.status_code == 200:
                    for p in search_resp.json().get("data", []):
                        if str(p.get("sku", "")).strip() == sku:
                            existing_id = p.get("id")
                            break

            if existing_id:
                update_payload = {k: v for k, v in payload.items() if k != "sku"}
                resp = await client.put(f"/products/{existing_id}", update_payload)
                action = "تحديث"
                product_id = existing_id
            else:
                resp = await client.post("/products", payload)
                action = "إنشاء"
                product_id = (
                    resp.json().get("data", {}).get("id")
                    if resp.status_code in (200, 201)
                    else None
                )

            if resp.status_code not in (200, 201):
                try:
                    err = resp.json().get("error", {}).get("message", resp.text)
                except Exception:
                    err = resp.text[:200]
                results.append(
                    {
                        "index": idx,
                        "name": name,
                        "success": False,
                        "error": f"فشل {action} ({resp.status_code}): {err}",
                    }
                )
                continue

            image_status = None
            image_url = item.get("image_url")
            if product_id and image_url:
                try:
                    local_path = download_image(str(product_id), image_url)
                    if local_path:
                        p = Path(local_path)
                        content = p.read_bytes()
                        up_resp = await client.upload_image(
                            product_id, p.name, content, "image/jpeg"
                        )
                        image_status = (
                            "✅ صورة مرفوعة"
                            if up_resp.status_code in (200, 201)
                            else f"⚠️ ({up_resp.status_code})"
                        )
                    else:
                        image_status = "⚠️ فشل تحميل الصورة محلياً"
                except Exception as e:
                    image_status = f"⚠️ {type(e).__name__}"

            results.append(
                {
                    "index": idx,
                    "name": name,
                    "product_id": product_id,
                    "action": action,
                    "success": True,
                    "image_status": image_status,
                }
            )

            await asyncio.sleep(0.4)

        except Exception as e:
            results.append(
                {
                    "index": idx,
                    "name": name,
                    "success": False,
                    "error": f"{type(e).__name__}: {str(e)[:200]}",
                }
            )

    success_count = sum(1 for r in results if r.get("success"))
    return {
        "results": results,
        "summary": {
            "total": len(results),
            "success": success_count,
            "failed": len(results) - success_count,
        },
    }


# ═══════════════════════════════════════════════════════════
# تصدير قالب Salla الرسمي (.xlsx)
# ═══════════════════════════════════════════════════════════
def export_salla_template(results: List[Dict[str, Any]], template_type: str = "new") -> bytes:
    if not results:
        raise ValueError("لا توجد نتائج للتصدير")

    if template_type == "update":
        template_path = BASE_DIR / "storage" / "salla_update_products_template.xlsx"
    else:
        template_path = BASE_DIR / "storage" / "salla_new_products_template.xlsx"

    if not template_path.exists():
        raise ValueError(f"قالب Salla غير موجود في المسار: {template_path}")

    wb = openpyxl.load_workbook(template_path)
    ws = wb.worksheets[0]

    valid_categories = set()
    valid_types = set()
    if len(wb.worksheets) > 1:
        ref_ws = wb.worksheets[1]
        for row_idx in range(2, ref_ws.max_row + 1):
            cat = ref_ws.cell(row=row_idx, column=1).value
            typ = ref_ws.cell(row=row_idx, column=2).value
            if cat and str(cat).strip():
                valid_categories.add(str(cat).strip())
            if typ and str(typ).strip():
                valid_types.add(str(typ).strip())

    headers = {}
    for col_idx in range(1, ws.max_column + 1):
        h = ws.cell(row=2, column=col_idx).value
        if h:
            headers[str(h).strip()] = col_idx

    DATA_START_ROW = 3

    def clean_number(v):
        if v is None or v == "":
            return ""
        s = str(v).strip()
        if s.endswith(".0"):
            s = s[:-2]
        return s

    def clean_price(v):
        if v is None or v == "":
            return ""
        try:
            return float(str(v).replace(",", "."))
        except (ValueError, TypeError):
            return ""

    def find_valid_category(user_cat):
        if not user_cat:
            return ""
        user_cat = str(user_cat).strip()
        if user_cat in valid_categories:
            return user_cat
        if "," in user_cat:
            first = user_cat.split(",")[0].strip()
            if first in valid_categories:
                return first
        if ">" in user_cat:
            main = user_cat.split(">")[0].strip()
            if main in valid_categories:
                return main
        for cat in valid_categories:
            if cat in user_cat or user_cat in cat:
                return cat
        return ""

    def find_valid_type(user_type):
        if not valid_types:
            return user_type or "منتج جاهز"
        if not user_type:
            for c in ["منتج جاهز", "أكل", "منتج رقمي"]:
                if c in valid_types:
                    return c
            return list(valid_types)[0]
        user_type = str(user_type).strip()
        if user_type in valid_types:
            return user_type
        for c in ["منتج جاهز", "أكل", "منتج رقمي", "خدمة حسب الطلب"]:
            if c in valid_types:
                return c
        return list(valid_types)[0]

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            ws.cell(row=row_idx, column=col_idx).value = None

    for i, r in enumerate(results):
        row_num = DATA_START_ROW + i

        # استخراج رابط الصورة مع التحقق والتشذيب
        image_url = (r.get("image_url") or "").strip()
        if not image_url and r.get("candidates") and len(r["candidates"]) > 0:
            image_url = (r["candidates"][0].get("image_url") or "").strip()
        if not image_url:
            image_url = (r.get("original_image_url") or "").strip()

        barcode = clean_number(r.get("barcode"))
        sku = clean_number(r.get("sku")) or ""

        field_map = {
            "النوع": "منتج",
            "أسم المنتج": r.get("name") or "",
            "تصنيف المنتج": find_valid_category(r.get("category")),
            "صورة المنتج": image_url,
            "وصف صورة المنتج": "",
            "نوع المنتج": find_valid_type(r.get("product_type")),
            "سعر المنتج": clean_price(r.get("price")),
            "الوصف": r.get("description") or "",
            "هل يتطلب شحن؟": "نعم",
            "رمز المنتج sku": sku,
            "سعر التكلفة": "",
            "السعر المخفض": "",
            "اقصي كمية لكل عميل": "",
            "إخفاء خيار تحديد الكمية": "",
            "اضافة صورة عند الطلب": "لا",
            "الوزن": r.get("weight") or "1",
            "وحدة الوزن": r.get("weight_unit") or "kg",
            "الماركة": "",
            "العنوان الترويجي": "",
            "تثبيت المنتج": "لا",
            "الباركود": barcode,
            "السعرات الحرارية": "",
            "MPN": "",
            "GTIN": "",
            "خاضع للضريبة ؟": "نعم",
            "سبب عدم الخضوع للضريبة": "",
        }

        if template_type == "update":
            product_no = r.get("no") or r.get("product_no") or r.get("row")
            if product_no:
                field_map["No."] = clean_number(product_no)

        for field_name, value in field_map.items():
            col_idx = headers.get(field_name.strip())
            if col_idx:
                ws.cell(row=row_num, column=col_idx).value = value

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ═══════════════════════════════════════════════════════════
# دوال مساعدة خاصة بقراءة الأعمدة وتنسيق البيانات
# ═══════════════════════════════════════════════════════════
def _safe_float(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except (ValueError, TypeError):
        return None


def _safe_int(v):
    if v is None or v == "":
        return None
    try:
        return int(float(str(v)))
    except (ValueError, TypeError):
        return None


def _clean_str(v):
    if v is None or v == "":
        return ""
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def _extract_weight_from_name(name: str) -> Tuple[float, str]:
    if not name:
        return (1.0, "kg")

    patterns = [
        (r"(\d+(?:[.,]\d+)?)\s*(?:كيلو|كجم|كيلوجرام|kg|kilogram)", "kg", 1),
        (r"(\d+(?:[.,]\d+)?)\s*(?:جرام|غرام|جم|غ|g|gram)", "kg", 0.001),
        (r"(\d+(?:[.,]\d+)?)\s*(?:ملليجرام|mg|milligram)", "kg", 0.000001),
        (r"(\d+(?:[.,]\d+)?)\s*(?:لتر|ل|liter|litre|l)", "kg", 1.0),
        (r"(\d+(?:[.,]\d+)?)\s*(?:مللي\s*لتر|مل|ml|milliliter)", "kg", 0.001),
    ]

    for pattern, unit, factor in patterns:
        match = re.search(pattern, name, re.IGNORECASE)
        if match:
            try:
                value = float(match.group(1).replace(",", "."))
                converted = value * factor
                if converted < 0.1:
                    return (round(converted * 1000, 2), "g")
                return (round(converted, 3), "kg")
            except (ValueError, TypeError):
                continue
    return (1.0, "kg")


def _read_excel_file(content: bytes, filename: str) -> pd.DataFrame:
    if filename.endswith(".csv"):
        for enc in ["utf-8-sig", "utf-8", "cp1256", "latin-1"]:
            try:
                raw_df = pd.read_csv(io.BytesIO(content), encoding=enc, header=None)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("فشل قراءة CSV بأي ترميز")
    else:
        engine = "openpyxl" if filename.endswith(".xlsx") else "xlrd"
        raw_df = pd.read_excel(io.BytesIO(content), engine=engine, header=None)

    header_row = _detect_header_row(raw_df)
    if header_row is None:
        raise ValueError("لم يتم العثور على صف العناوين في الملف")

    if filename.endswith(".csv"):
        for enc in ["utf-8-sig", "utf-8", "cp1256", "latin-1"]:
            try:
                df = pd.read_csv(
                    io.BytesIO(content), encoding=enc, header=header_row, dtype=str
                )
                break
            except UnicodeDecodeError:
                continue
    else:
        engine = "openpyxl" if filename.endswith(".xlsx") else "xlrd"
        df = pd.read_excel(
            io.BytesIO(content), engine=engine, header=header_row, dtype=str
        )

    df = df.dropna(axis=1, how="all")
    df = df.dropna(how="all")
    df = _drop_total_rows(df)
    df.columns = [str(c).strip() for c in df.columns]

    name_col = _detect_columns(df).get("name")
    if name_col:
        df = df[df[name_col].notna()]
        df = df[df[name_col].astype(str).str.strip() != ""]

    return df


def _detect_header_row(df: pd.DataFrame) -> Optional[int]:
    keywords = [
        "أسم المنتج",
        "اسم المنتج",
        "رمز المنتج sku",
        "سعر المنتج",
        "صورة المنتج",
        "تصنيف المنتج",
        "حالة المنتج",
        "الباركود",
        "name",
        "barcode",
        "price",
        "quantity",
        "sku",
        "اسم",
        "باركود",
        "سعر",
    ]

    for idx in range(min(10, len(df))):
        row = df.iloc[idx]
        matches = 0
        for cell in row:
            if pd.isna(cell):
                continue
            cell_text = str(cell).strip().lower()
            if any(kw.lower() in cell_text for kw in keywords):
                matches += 1
        if matches >= 3:
            return idx
    return None


def _drop_total_rows(df: pd.DataFrame) -> pd.DataFrame:
    total_markers = [
        "الإجمالي",
        "اجمالي",
        "إجمالي",
        "المجموع",
        "مجموع",
        "total",
        "sum",
        "grand total",
        "subtotal",
        "ملاحظات",
        "notes",
        "تعليقات",
    ]

    def is_total_row(row) -> bool:
        for cell in row:
            if pd.isna(cell):
                continue
            cell_text = str(cell).strip().lower()
            if any(marker in cell_text for marker in total_markers):
                return True
        return False

    mask = df.apply(is_total_row, axis=1)
    return df[~mask]


def _detect_columns(df: pd.DataFrame) -> Dict[str, str]:
    mapping = {}
    aliases = {
        "no": ["No.", "الرقم", "م", "id"],
        "name": [
            "اسم المنتج الغذائي",
            "أسم المنتج",
            "اسم المنتج",
            "الاسم",
            "اسم",
            "name",
            "title",
            "المنتج",
        ],
        "barcode": ["الباركود", "باركود", "barcode", "GTIN", "EAN", "UPC"],
        "sku": ["رمز المنتج", "SKU", "sku", "الرمز"],
        "price": ["سعر المنتج", "السعر", "سعر", "price"],
        "sale_price": ["السعر المخفض"],
        "quantity": ["الكمية", "quantity", "qty", "stock", "المخزون"],
        "description": ["الوصف", "description", "وصف"],
        "image_url": ["صورة المنتج", "صورة", "image"],
        "image_alt": ["وصف صورة المنتج"],
        "category": ["تصنيف المنتج", "التصنيف", "الفئة", "category"],
        "type": ["النوع"],
        "product_type": ["نوع المنتج"],
        "status": ["حالة المنتج", "الحالة", "status"],
        "weight": ["الوزن", "weight"],
        "weight_unit": ["وحدة الوزن"],
        "brand": ["الماركة", "brand"],
        "cost": ["سعر التكلفة"],
        "requires_shipping": ["هل يتطلب شحن"],
        "taxable": ["خاضع للضريبة"],
        "calories": ["السعرات الحرارية"],
        "mpn": ["MPN"],
        "promo_title": ["العنوان الترويجي"],
        "pin": ["تثبيت المنتج"],
        "max_per_customer": ["اقصي كمية لكل عميل"],
        "hide_qty": ["إخفاء خيار تحديد الكمية"],
        "image_on_order": ["اضافة صورة عند الطلب"],
    }

    cols_normalized = []
    for c in df.columns:
        cleaned = str(c).strip()
        no_parens = re.sub(r"\s*[\(\[][^)\]]*[\)\]]\s*", "", cleaned).strip()
        cols_normalized.append(
            {
                "original": c,
                "lower": cleaned.lower(),
                "no_parens_lower": no_parens.lower(),
            }
        )

    for key, names in aliases.items():
        sorted_names = sorted(names, key=len, reverse=True)
        matched = False

        for alias in sorted_names:
            alias_lower = alias.lower()
            for col in cols_normalized:
                if col["no_parens_lower"] == alias_lower or col["lower"] == alias_lower:
                    mapping[key] = col["original"]
                    matched = True
                    break
            if matched:
                break

        if matched:
            continue

        for alias in sorted_names:
            alias_lower = alias.lower()
            if len(alias_lower) < 3 and alias_lower not in ("رقم", "sku"):
                continue
            for col in cols_normalized:
                if alias_lower in col["no_parens_lower"] or alias_lower in col["lower"]:
                    mapping[key] = col["original"]
                    matched = True
                    break
            if matched:
                break

    return mapping


def _get_cell(row, column) -> str:
    if not column:
        return ""
    value = row.get(column)
    if pd.isna(value):
        return ""

    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]

    if "e+" in s.lower() or "e-" in s.lower():
        try:
            s = str(int(float(s)))
        except (ValueError, OverflowError):
            pass

    return s
